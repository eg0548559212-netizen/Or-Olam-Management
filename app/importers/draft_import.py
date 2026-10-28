"""ייבוא טיוטת העץ מקובץ docx (D-009) ל-draft, draft_node ו-draft_note (06-planning.md).

הרצה:
    python3 -m app.importers.draft_import data/draft/draft-2026-10-21.docx --label "הצעת סידור פנימי 21/10"

כללי הקריאה של הטיוטה מ-21.10:
- כותרת "שלוחה N: שם" היא צומת ברמה הראשונה תחת התפריט הראשי (1.N). הפסקה שאחריה היא תיאור.
- רמת הרשימה ב-Word קובעת את ההיררכיה.
- המספר בטקסט כתוב הפוך ("2.6.2.2." = 1.2.2.6.2). הספרה הראשונה היא הספרה המקומית.
  הנתיב המלא מחושב מהאב. הנתיב הכתוב נשמר ב-hints.stated_path לבדיקה.
- צבע: שחור קיימת, מג'נטה חדשה, כחול בנייה מיוחדת, אדום שאלה, כתום הערת בנייה.
- שורה בלי מספר שכולה אדום או כתום היא הערה על האב. אחרת היא צומת בלי מספר.
"""
import argparse
import hashlib
import json
import os
import re
from collections import Counter

import docx

from app.db.connect import connect, now

COLOR_STATUS = {None: "existing", "FF00FF": "new", "0000FF": "new_special_build"}
NOTE_COLORS = {"FF0000": "question", "FF9900": "build_note"}

SECTION_RE = re.compile(r"^שלוחה\s+(\d+)\s*:\s*(.*)$")
NUMBER_RE = re.compile(r"^\s*((?:\d+\.)+)\s*")
GLOBAL_BLOCKS = {
    "בכל הקו": ("global", None),
    "שאלות שנותרו": ("open_question", "question"),
    "דרושים מוסרים": ("global", "comment"),
}


def run_color(run):
    c = run.font.color
    if c is None or c.type is None or c.rgb is None:
        return None
    rgb = str(c.rgb).upper()
    return None if rgb == "000000" else rgb


def list_level(p):
    ppr = p._p.pPr
    if ppr is None or ppr.numPr is None or ppr.numPr.ilvl is None:
        return None
    return int(ppr.numPr.ilvl.val)


def split_runs(p):
    """מחלק פסקה לחלק הכותרת ולחלקי ההערות לפי צבע. מחזיר (title_runs, notes)."""
    title, notes = [], []
    for r in p.runs:
        if not r.text:
            continue
        color = run_color(r)
        if color in NOTE_COLORS:
            if notes and notes[-1][0] == NOTE_COLORS[color]:
                notes[-1][1].append(r.text)
            else:
                notes.append((NOTE_COLORS[color], [r.text]))
        else:
            title.append((color, r.text))
    return title, [(k, clean("".join(t))) for k, t in notes if clean("".join(t))]


def clean(text):
    return re.sub(r"\s+", " ", text).strip(" -–:‏‎\t")


def title_status(title_runs):
    """הסטטוס לפי הצבע הדומיננטי של הכותרת, בלי המספר."""
    weights = Counter()
    for color, text in title_runs:
        stripped = NUMBER_RE.sub("", text) if not weights else text
        n = len(stripped.strip())
        if n:
            weights[color if color in COLOR_STATUS else None] += n
    if not weights:
        return "existing", None
    color = weights.most_common(1)[0][0]
    if weights.get("FF00FF") and color is None:
        color = "FF00FF"  # מספר בשחור ושם במג'נטה
    if weights.get("0000FF") and color is None:
        color = "0000FF"
    return COLOR_STATUS.get(color, "unknown"), color


HINT_PATTERNS = [
    ("episodes", re.compile(r"(\d+)\s*(?:פרקים|שיעורים)")),
    ("closed_series", re.compile(r"סדרה סגורה")),
    ("renewal", re.compile(r"מתחדשת\s+(?:מדי|בכל)\s*([^\s?.,]*)")),
    ("copy_to", re.compile(r"יועתקו\s+גם\s+ל-?\s*(\d+)")),
    ("redirect_to", re.compile(r"(?:מועבר|מעבר ישיר)\s+לשלוחה\s+(\d+)")),
]
LECTURER_RE = re.compile(r"\(([^()0-9?]{2,40})\)")


def extract_hints(text):
    hints = {}
    for name, rx in HINT_PATTERNS:
        m = rx.search(text)
        if m:
            hints[name] = m.group(1) if m.groups() else True
    lecturers = [m.strip() for m in LECTURER_RE.findall(text) if len(m.split()) <= 3]
    if lecturers:
        hints["lecturers_text"] = lecturers
    return hints


def parse(docx_path):
    """מחזיר (nodes, notes). node: dict עם id, parent, ..., notes: (node_id|None, scope, kind, text, color)."""
    doc = docx.Document(docx_path)
    nodes, notes = [], []
    stack = {}  # רמת רשימה → node
    section = None
    global_block = None
    expect_description = False
    position = 0

    def add_node(parent, digit, raw_number, raw_text, title, status, color, span, kind="extension", hints=None):
        nonlocal position
        position += 1
        node = {
            "node_id": len(nodes) + 1,
            "parent": parent,
            "position": position,
            "digit": digit,
            "raw_number_text": raw_number,
            "raw_text": raw_text,
            "title": title,
            "node_status": status,
            "node_kind": kind,
            "hints": hints or {},
            "source_color": color,
            "source_span": span,
        }
        parent_path = parent["full_path"] if parent else None
        if digit is None or (parent is not None and parent_path is None):
            node["full_path"] = None
        else:
            node["full_path"] = f"{parent_path}.{digit}" if parent else digit
        nodes.append(node)
        return node

    root = add_node(None, "1", None, None, "תפריט ראשי", "existing", None, {"paragraph": None}, kind="heading")

    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue
        span = {"paragraph": i}
        level = list_level(p)
        title_runs, inline_notes = split_runs(p)

        m = SECTION_RE.match(text) if level is None else None
        if m:
            own = SECTION_RE.match(clean("".join(t for _, t in title_runs)))
            title_text = clean((own or m).group(2))
            status, color = title_status([(c, t) for c, t in title_runs])
            section = add_node(root, m.group(1), m.group(1), text, title_text, status, color, span)
            for kind, ntext in inline_notes:
                notes.append((section["node_id"], "node", kind, ntext, None))
            stack = {}
            global_block = None
            expect_description = True
            continue

        if level is None:
            block = next((k for k in GLOBAL_BLOCKS if text.startswith(k)), None)
            if block:
                global_block = block
                section = None
                expect_description = False
                continue
            if section is not None and expect_description:
                section["hints"]["description"] = clean("".join(t for _, t in title_runs))
                for kind, ntext in inline_notes:
                    notes.append((section["node_id"], "node", kind, ntext, None))
                expect_description = False
            continue
        expect_description = False

        # ── פסקת רשימה ──
        if section is None:
            if global_block is None:
                continue  # מקרא והוראות בראש המסמך
            scope, forced_kind = GLOBAL_BLOCKS[global_block]
            color_kind = Counter(NOTE_COLORS.get(run_color(r)) for r in p.runs if r.text.strip()).most_common(1)[0][0]
            kind = forced_kind or color_kind or "comment"
            prefix = f"{global_block}: " if global_block == "דרושים מוסרים" else ""
            notes.append((None, scope, kind, prefix + clean(text), None))
            continue

        parent = stack.get(level - 1, section) if level > 0 else section
        # המספר נקרא מהטקסט המלא: לפעמים ספרה בודדת צבועה בצבע של הערה.
        num = NUMBER_RE.match(text)
        inline_notes = [(k, t) for k, t in inline_notes if not re.fullmatch(r"[\d.\s]*", t)]
        title_text = clean(re.sub(r"\(\s*\)", "", re.sub(r"^[\d.]+\s*", "", "".join(t for _, t in title_runs).strip())))

        if not num and not title_text and inline_notes:
            for kind, ntext in inline_notes:
                notes.append((parent["node_id"], "node", kind, ntext, None))
            parent["hints"].update({k: v for k, v in extract_hints(text).items() if k not in parent["hints"]})
            continue

        digit, raw_number, hints = None, None, extract_hints(text)
        if num:
            raw_number = num.group(1)
            segs = [s for s in raw_number.split(".") if s]
            digit = segs[0]
            if len(segs) > 1:
                hints["stated_path"] = "1." + ".".join(reversed(segs))
        status, color = title_status(title_runs)
        kind = "extension"
        if not title_text or re.fullmatch(r"[?_\s]+", title_text):
            kind = "placeholder"
        node = add_node(parent, digit, raw_number, text, title_text or None, status, color, span, kind=kind, hints=hints)
        for kind_, ntext in inline_notes:
            notes.append((node["node_id"], "node", kind_, ntext, None))
        stack[level] = node
        for deeper in [lv for lv in stack if lv > level]:
            del stack[deeper]

    return nodes, notes


def run(docx_path, db_file, label, plan_name="תכנית מעבר לעץ החדש"):
    nodes, notes = parse(docx_path)
    db = connect(db_file)
    ts = now()

    plan = db.execute("SELECT plan_id FROM migration_plan ORDER BY plan_id LIMIT 1").fetchone()
    if plan:
        plan_id = plan["plan_id"]
    else:
        plan_id = db.execute(
            "INSERT INTO migration_plan (name, status, created_by, created_at) VALUES (?, 'drafting', 'import', ?)",
            (plan_name, ts),
        ).lastrowid
        take_base_snapshot(db, plan_id, ts)

    data = open(docx_path, "rb").read()
    file_id = db.execute(
        """INSERT INTO uploaded_file (filename, format, uploaded_by, uploaded_at, sha256, stored_path, parse_status, parse_log)
           VALUES (?, 'docx', 'import', ?, ?, ?, 'ok', ?)""",
        (os.path.basename(docx_path), ts, hashlib.sha256(data).hexdigest(), docx_path,
         json.dumps({"nodes": len(nodes), "notes": len(notes)}, ensure_ascii=False)),
    ).lastrowid

    # הקובץ הוא גרסה מסוג imported. אם אין עדיין טיוטת עבודה, היא נוצרת ממנו (D-108).
    imported_id = insert_draft(db, plan_id, "imported", file_id, None, label, ts, nodes, notes)
    working = db.execute("SELECT draft_id FROM draft WHERE plan_id = ? AND kind = 'working'", (plan_id,)).fetchone()
    working_id = working["draft_id"] if working else insert_draft(
        db, plan_id, "working", file_id, imported_id, f"טיוטת עבודה (מתוך {label})", ts, nodes, notes)
    db.commit()
    return {
        "plan_id": plan_id,
        "imported_draft_id": imported_id,
        "working_draft_id": working_id,
        "nodes": len(nodes),
        "notes": len(notes),
        "by_status": dict(Counter(n["node_status"] for n in nodes)),
        "unnumbered": sum(1 for n in nodes if n["digit"] is None),
        "placeholders": sum(1 for n in nodes if n["node_kind"] == "placeholder"),
        "stated_path_mismatch": sum(
            1 for n in nodes if n["hints"].get("stated_path") and n["full_path"] and n["hints"]["stated_path"] != n["full_path"]),
    }


def insert_draft(db, plan_id, kind, file_id, based_on, label, ts, nodes, notes):
    draft_id = db.execute(
        """INSERT INTO draft (plan_id, kind, source_file_id, based_on_draft_id, label, created_by, created_at)
           VALUES (?, ?, ?, ?, ?, 'import', ?)""",
        (plan_id, kind, file_id, based_on, label, ts),
    ).lastrowid
    for n in nodes:
        db.execute(
            """INSERT INTO draft_node (draft_id, node_id, parent_node_id, position, digit, full_path, raw_number_text,
                 raw_text, title, node_status, node_kind, planned_cls_type, hints, source_color, source_span)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,NULL,?,?,?)""",
            (draft_id, n["node_id"], n["parent"]["node_id"] if n["parent"] else None, n["position"], n["digit"],
             n["full_path"], n["raw_number_text"], n["raw_text"], n["title"], n["node_status"], n["node_kind"],
             json.dumps(n["hints"], ensure_ascii=False) if n["hints"] else None, n["source_color"],
             json.dumps(n["source_span"])),
        )
    for node_id, scope, kind, text, color in notes:
        db.execute(
            """INSERT INTO draft_note (draft_id, node_id, scope, note_kind, text, source_color, status, author, created_at)
               VALUES (?, ?, ?, ?, ?, ?, 'open', 'import', ?)""",
            (draft_id, node_id, scope, kind, text,
             color or {"question": "FF0000", "build_note": "FF9900"}.get(kind), ts),
        )
    return draft_id


def take_base_snapshot(db, plan_id, ts):
    run = db.execute("SELECT MAX(run_id) FROM tree_sync_run WHERE status = 'ok'").fetchone()[0]
    snap_id = db.execute(
        "INSERT INTO base_tree_snapshot (plan_id, taken_at, source_run_id, note) VALUES (?, ?, ?, 'צילום בתחילת התכנית')",
        (plan_id, ts, run),
    ).lastrowid
    db.execute(
        """INSERT INTO base_tree_snapshot_node (snapshot_id, sclass_id, path, name, cls_type, parent_sclass_id)
           SELECT ?, sclass_id, path, name, cls_type, parent_sclass_id FROM extension WHERE missing_since IS NULL""",
        (snap_id,),
    )
    return snap_id


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("docx_path")
    ap.add_argument("--db", default="data/or-olam.db")
    ap.add_argument("--label", default="הצעת סידור פנימי 21/10")
    args = ap.parse_args()
    print(json.dumps(run(args.docx_path, args.db, args.label), ensure_ascii=False, indent=2))
